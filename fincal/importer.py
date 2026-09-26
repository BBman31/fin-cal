"""Convert the original budget workbook (template layout) into settings + a slim log."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.worksheet.worksheet import Worksheet

from fincal.mapping import Lookup, normalize
from fincal.models import BUCKETS, FixedCost, FlexCategory, OpenClose, Settings
from fincal.spending import SHEET_NAME, clean_log


def _num(value) -> float | None:
    if value is None or isinstance(value, str) and not value.strip():
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _sheet(wb, name: str) -> Worksheet | None:
    for ws in wb.worksheets:
        if normalize(ws.title) == normalize(name):
            return ws
    return None


def _find(ws: Worksheet, label: str) -> tuple[int, int] | None:
    """(row, col) of the first cell whose text starts with ``label`` (case-insensitive)."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and normalize(cell.value).startswith(normalize(label)):
                return cell.row, cell.column
    return None


def _table_below(ws: Worksheet, header: tuple[int, int], width: int):
    """Yield rows (lists of ``width`` values) under a header cell until the first blank key."""
    row, col = header
    for r in range(row + 1, ws.max_row + 1):
        values = [ws.cell(r, col + i).value for i in range(width)]
        if values[0] is None or _text(values[0]) == "":
            return
        yield values


def read_overview(ws: Worksheet, settings: Settings) -> None:
    pos = _find(ws, "Net Monthly Income")
    if pos:
        settings.income = _num(ws.cell(pos[0], pos[1] + 1).value) or 0.0
    weights = {}
    header = _find(ws, "Budget split")
    if header:
        for label, weight, *_ in _table_below(ws, header, 2):
            bucket = next((b for b in BUCKETS if normalize(label).startswith(b.lower())), None)
            if bucket:
                weights[bucket] = _num(weight) or 0.0
    if weights:
        settings.bucket_weights = {b: weights.get(b, 0.0) for b in BUCKETS}


def read_fixed(ws: Worksheet, settings: Settings) -> None:
    header = _find(ws, "Category")
    if not header:
        return
    for category, amount, comment in _table_below(ws, header, 3):
        if normalize(category) in ("total", "remaining"):
            break
        settings.fixed_costs.append(FixedCost(_text(category), _num(amount) or 0.0, _text(comment)))


def read_flexible(ws: Worksheet, settings: Settings) -> None:
    header = _find(ws, "Category")
    if not header:
        return
    for category, weight in _table_below(ws, header, 2):
        settings.flexible_categories.append(FlexCategory(_text(category), _num(weight) or 0.0))


def read_open_close(ws: Worksheet, settings: Settings) -> None:
    header = _find(ws, "Date")
    if not header:
        return
    for when, open_, close in _table_below(ws, header, 3):
        if isinstance(when, datetime | date):
            month = when.strftime("%Y-%m")
        else:
            parsed = pd.to_datetime(_text(when), errors="coerce", dayfirst=True)
            if pd.isna(parsed):
                continue
            month = parsed.strftime("%Y-%m")
        settings.open_close.append(OpenClose(month, _num(open_), _num(close)))


def read_mappings(ws: Worksheet, settings: Settings) -> list[str]:
    """Find the 'Category | Bucket' and 'Item | Category' header pairs anywhere on the sheet."""
    warnings = []
    for row in ws.iter_rows():
        for cell in row:
            left = normalize(cell.value)
            right = normalize(ws.cell(cell.row, cell.column + 1).value)
            if (left, right) == ("category", "bucket"):
                target, name = "category_buckets", "Category→Bucket"
            elif (left, right) == ("item", "category"):
                target, name = "item_categories", "Item→Category"
            else:
                continue
            pairs = [(_text(k), _text(v)) for k, v in _table_below(ws, (cell.row, cell.column), 2)]
            lookup = Lookup.from_pairs(pairs)
            for key, values in lookup.conflicts.items():
                warnings.append(f"{name}: '{key}' maps to {sorted(values)}; kept the first")
            mapping = getattr(settings, target)
            seen = set()
            for key, value in pairs:
                if key and value and normalize(key) not in seen:
                    seen.add(normalize(key))
                    mapping[key] = lookup.exact[normalize(key)]
    return warnings


def read_settings(path: Path) -> tuple[Settings, list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    settings = Settings(bucket_weights={})
    warnings: list[str] = []
    readers = [
        ("Overview", read_overview),
        ("Fixed", read_fixed),
        ("Flexible Budget", read_flexible),
        ("OpenClose", read_open_close),
    ]
    for name, reader in readers:
        ws = _sheet(wb, name)
        if ws is None:
            warnings.append(f"Sheet '{name}' not found; skipped")
        else:
            reader(ws, settings)
    ws = _sheet(wb, "Mapping tables")
    if ws is None:
        warnings.append("Sheet 'Mapping tables' not found; skipped")
    else:
        warnings += read_mappings(ws, settings)
    if not settings.bucket_weights:
        settings.bucket_weights = Settings().bucket_weights
    warnings += settings.validate()
    return settings, warnings


def read_log(path: Path, settings: Settings) -> pd.DataFrame:
    """Slim log: Date, Item, Amount, Category (only where it differs from the mapping), Comment."""
    raw = pd.read_excel(path, sheet_name=None)
    sheet = next((df for name, df in raw.items() if normalize(name) == normalize(SHEET_NAME)), None)
    if sheet is None:
        return pd.DataFrame(columns=["Date", "Item", "Amount", "Category", "Comment"])
    log = clean_log(sheet)
    items = Lookup.from_dict(settings.item_categories)
    derived = log["item"].map(lambda i: items.get(i))
    keep = log["category_override"].notna() & (
        log["category_override"].map(normalize) != derived.map(normalize)
    )
    return pd.DataFrame(
        {
            "Date": log["date"].dt.date,
            "Item": log["item"],
            "Amount": log["amount"],
            "Category": log["category_override"].where(keep, None),
            "Comment": log["comment"],
        }
    )


def write_log(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl", date_format="YYYY-MM-DD") as writer:
        df.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        ws = writer.sheets[SHEET_NAME]
        for col, width in zip("ABCDE", (12, 32, 12, 18, 40), strict=True):
            ws.column_dimensions[col].width = width
        for cell in ws["A"][1:]:
            cell.number_format = "YYYY-MM-DD"
        ws.freeze_panes = "A2"
